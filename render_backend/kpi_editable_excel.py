from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


BG = "EFEFEF"
NAVY = "0B2F66"
TEAL = "13A6A8"
PINK = "F8DADB"
MAGENTA = "E91E63"
LIGHT_GREY = "E1E4E8"
GRID = "CFCFCF"
TEXT = "5F6368"
DARK = "24344D"
RED = "E60039"
WHITE = "FFFFFF"


def _num(value: Any, default: float = 0.0) -> float:
    try:
        if value is None or value == "":
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def _setting(settings: dict[str, Any], key: str, default: float) -> float:
    return _num(settings.get(key), default)


def _pct(value: Any) -> float:
    return _num(value) / 100.0


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _side(color: str = GRID, style: str = "thin") -> Side:
    return Side(style=style, color=color)


def _border(color: str = GRID, style: str = "thin") -> Border:
    side = _side(color, style)
    return Border(left=side, right=side, top=side, bottom=side)


def _font(size: float = 9, bold: bool = False, color: str = DARK) -> Font:
    return Font(name="Segoe UI", size=size, bold=bold, color=color)


def _align(horizontal: str = "center", vertical: str = "center", wrap: bool = False) -> Alignment:
    return Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap)


def _style_range(
    ws,
    row1: int,
    col1: int,
    row2: int,
    col2: int,
    *,
    fill: str | None = None,
    border: Border | None = None,
) -> None:
    for row in ws.iter_rows(min_row=row1, min_col=col1, max_row=row2, max_col=col2):
        for cell in row:
            if fill:
                cell.fill = _fill(fill)
            if border:
                cell.border = border


def _merge_value(ws, cell_range: str, value: Any, *, font: Font, fill: str, border: Border, number_format: str | None = None) -> None:
    ws.merge_cells(cell_range)
    cell = ws[cell_range.split(":")[0]]
    cell.value = value
    cell.font = font
    cell.fill = _fill(fill)
    cell.border = border
    cell.alignment = _align()
    if number_format:
        cell.number_format = number_format
    for row in ws[cell_range]:
        for item in row:
            item.fill = _fill(fill)
            item.border = border
            item.alignment = _align()


def _write_bar(ws, row: int, start_col: int, end_col: int, formula_value: str, color: str) -> None:
    steps = max(end_col - start_col + 1, 1)
    active_fill = _fill(color)
    empty_fill = _fill(LIGHT_GREY)
    for offset, col in enumerate(range(start_col, end_col + 1), start=1):
        cell = ws.cell(row=row, column=col)
        cell.value = f'=IF({formula_value}>={offset / steps:.6f},1,"")'
        cell.number_format = ";;;"
        cell.fill = empty_fill
        cell.border = Border()
    rng = f"{get_column_letter(start_col)}{row}:{get_column_letter(end_col)}{row}"
    ws.conditional_formatting.add(
        rng,
        CellIsRule(operator="greaterThan", formula=["0"], fill=active_fill),
    )


def _write_metric_pair(
    ws,
    top_row: int,
    bottom_row: int,
    left_cols: tuple[int, int],
    right_cols: tuple[int, int],
    left_title: str,
    left_value: str,
    left_bar: str,
    left_footer: str,
    left_tone: str,
    right_title: str,
    right_value: str,
    right_bar: str,
    right_footer: str,
    right_tone: str,
    *,
    left_number_format: str = "0.0%",
    right_number_format: str = "0.0%",
) -> None:
    left_fill = PINK if left_tone == "pink" else WHITE
    right_fill = PINK if right_tone == "pink" else WHITE
    left_bar_color = MAGENTA if left_tone == "pink" else TEAL
    right_bar_color = MAGENTA if right_tone == "pink" else TEAL
    l1, l2 = left_cols
    r1, r2 = right_cols
    _merge_value(ws, f"{get_column_letter(l1)}{top_row}:{get_column_letter(l2)}{top_row}", left_title, font=_font(11, True, TEXT), fill=left_fill, border=_border())
    _merge_value(ws, f"{get_column_letter(l1)}{top_row + 1}:{get_column_letter(l2)}{top_row + 3}", left_value, font=_font(20, True, TEXT), fill=left_fill, border=_border(), number_format=left_number_format)
    _write_bar(ws, top_row + 4, l1 + 1, l2 - 1, left_bar, left_bar_color)
    _merge_value(ws, f"{get_column_letter(l1)}{top_row + 5}:{get_column_letter(l2)}{top_row + 5}", left_footer, font=_font(8, False, "66707A"), fill=WHITE, border=_border(), number_format="0.0%")

    _merge_value(ws, f"{get_column_letter(r1)}{top_row}:{get_column_letter(r2)}{top_row}", right_title, font=_font(11, True, TEXT), fill=right_fill, border=_border())
    _merge_value(ws, f"{get_column_letter(r1)}{top_row + 1}:{get_column_letter(r2)}{top_row + 3}", right_value, font=_font(20, True, TEXT), fill=right_fill, border=_border(), number_format=right_number_format)
    _write_bar(ws, top_row + 4, r1 + 1, r2 - 1, right_bar, right_bar_color)
    _merge_value(ws, f"{get_column_letter(r1)}{top_row + 5}:{get_column_letter(r2)}{top_row + 5}", right_footer, font=_font(8, False, "66707A"), fill=WHITE, border=_border(), number_format="0.0%")
    _merge_value(ws, f"{get_column_letter(l1)}{top_row + 6}:{get_column_letter(r2)}{bottom_row}", "", font=_font(), fill=WHITE, border=_border())


def _group_title(group: str) -> tuple[str, str]:
    text = str(group or "").upper()
    if "REZAG" in text:
        return "REZAGADO", "Equipos de Rezagado"
    return "BARRENACION", "Equipos de Barrenacion"


def _configure_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A35"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.22
    ws.page_margins.right = 0.22
    ws.page_margins.top = 0.28
    ws.page_margins.bottom = 0.28
    ws.print_area = "A1:AD46"
    ws.sheet_properties.tabColor = NAVY
    for col in range(1, 31):
        ws.column_dimensions[get_column_letter(col)].width = 4.4
    ws.column_dimensions["I"].width = 1.2
    for col in range(10, 23):
        ws.column_dimensions[get_column_letter(col)].width = 4.7
    for col in range(32, 41):
        ws.column_dimensions[get_column_letter(col)].hidden = True
    for row in range(1, 61):
        ws.row_dimensions[row].height = 20
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 8
    ws.row_dimensions[30].height = 28


def _write_table(ws, report: dict[str, Any], start_row: int, start_col: int, settings: dict[str, Any]) -> int:
    headers = ["# Eco", "Equipo", "Hrs Periodo", "Hrs MP", "Hrs MC", "Hrs Trab", "# Paradas", "% Disp", "% Util", "Confiabilidad", "TMEF", "TMPR", "Estatus"]
    widths = [7, 25, 11, 8, 8, 9, 9, 8, 8, 13, 8, 8, 14]
    for idx, width in enumerate(widths, start=start_col):
        ws.column_dimensions[get_column_letter(idx)].width = width
    header_fill = _fill(WHITE)
    header_font = _font(8, True, "404040")
    cell_font = _font(8, False, DARK)
    percent_font = _font(8, False, "00A6B2")
    red_font = _font(8, False, RED)
    table_border = _border("000000")
    for idx, header in enumerate(headers, start=start_col):
        cell = ws.cell(start_row, idx, header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = table_border
        cell.alignment = _align(wrap=True)
    rows = report.get("rows") if isinstance(report.get("rows"), list) else []
    for row_offset, row in enumerate(rows, start=1):
        excel_row = start_row + row_offset
        values = [
            row.get("code") or "",
            row.get("description") or "",
            _num(row.get("period")),
            _num(row.get("mp")),
            _num(row.get("mc")),
            _num(row.get("worked")),
            _num(row.get("stops")),
        ]
        for col_offset, value in enumerate(values, start=0):
            cell = ws.cell(excel_row, start_col + col_offset, value)
            cell.font = cell_font
            cell.fill = header_fill
            cell.border = table_border
            cell.alignment = _align()
            if isinstance(value, (int, float)):
                cell.number_format = "0.0"
        status = "FUERA" if row.get("out") else str(row.get("status") or "Disponible")
        period_cell = f"{get_column_letter(start_col + 2)}{excel_row}"
        mp_cell = f"{get_column_letter(start_col + 3)}{excel_row}"
        mc_cell = f"{get_column_letter(start_col + 4)}{excel_row}"
        worked_cell = f"{get_column_letter(start_col + 5)}{excel_row}"
        stops_cell = f"{get_column_letter(start_col + 6)}{excel_row}"
        status_cell = f"{get_column_letter(start_col + 12)}{excel_row}"
        formulas = [
            f'=IF(UPPER(${status_cell})="FUERA","FUERA",IF({period_cell}>0,MAX(MIN(({period_cell}-{mp_cell}-{mc_cell})/{period_cell},1),0),0))',
            f'=IF(UPPER(${status_cell})="FUERA","FUERA",IF(({period_cell}-{mp_cell}-{mc_cell})>0,MAX(MIN({worked_cell}/({period_cell}-{mp_cell}-{mc_cell}),1),0),0))',
            f'=IF(UPPER(${status_cell})="FUERA",0,IF({get_column_letter(start_col + 10)}{excel_row}>0,EXP(-($AF$6/{get_column_letter(start_col + 10)}{excel_row})),IF({period_cell}>0,1,0)))',
            f'=IF(UPPER(${status_cell})="FUERA",0,IF({stops_cell}>0,({period_cell}-{mp_cell}-{mc_cell})/{stops_cell},({period_cell}-{mp_cell}-{mc_cell})))',
            f'=IF(UPPER(${status_cell})="FUERA",0,IF({stops_cell}>0,{mc_cell}/{stops_cell},0))',
        ]
        for col_offset, formula in zip(range(7, 12), formulas):
            cell = ws.cell(excel_row, start_col + col_offset, formula)
            cell.font = red_font if col_offset in (7, 8) and row.get("out") else percent_font if col_offset in (7, 8, 9) else cell_font
            cell.fill = header_fill
            cell.border = table_border
            cell.alignment = _align()
            cell.number_format = "0.0%" if col_offset in (7, 8, 9) else "0.0"
        cell = ws.cell(excel_row, start_col + 12, status)
        cell.font = _font(8, True, RED if "FUERA" in status.upper() else "00A6B2" if status.upper() == "DISPONIBLE" else DARK)
        cell.fill = header_fill
        cell.border = table_border
        cell.alignment = _align()
    total_row = start_row + len(rows) + 1
    group = report.get("group") or ""
    total_label = "Total Equipos de Rezagado" if "REZAG" in str(group).upper() else "Total Equipos de Barrenacion"
    ws.cell(total_row, start_col, "")
    ws.cell(total_row, start_col + 1, total_label)
    for col in range(start_col, start_col + len(headers)):
        ws.cell(total_row, col).font = _font(8, True, DARK)
        ws.cell(total_row, col).fill = header_fill
        ws.cell(total_row, col).border = table_border
        ws.cell(total_row, col).alignment = _align()
    first_data = start_row + 1
    last_data = total_row - 1
    for col_offset in range(2, 7):
        col = get_column_letter(start_col + col_offset)
        ws.cell(total_row, start_col + col_offset, f"=SUM({col}{first_data}:{col}{last_data})")
        ws.cell(total_row, start_col + col_offset).number_format = "0.0"
    helper_first = 3
    helper_last = helper_first + len(rows) - 1
    ws.cell(total_row, start_col + 7, f"=AVERAGE(AJ{helper_first}:AJ{helper_last})/100" if rows else 0)
    ws.cell(total_row, start_col + 8, f"=AVERAGE(AK{helper_first}:AK{helper_last})/100" if rows else 0)
    ws.cell(total_row, start_col + 10, f'=IF({get_column_letter(start_col + 6)}{total_row}>0,({get_column_letter(start_col + 2)}{total_row}-{get_column_letter(start_col + 3)}{total_row}-{get_column_letter(start_col + 4)}{total_row})/{get_column_letter(start_col + 6)}{total_row},({get_column_letter(start_col + 2)}{total_row}-{get_column_letter(start_col + 3)}{total_row}-{get_column_letter(start_col + 4)}{total_row}))')
    ws.cell(total_row, start_col + 11, f"=IF({get_column_letter(start_col + 6)}{total_row}>0,{get_column_letter(start_col + 4)}{total_row}/{get_column_letter(start_col + 6)}{total_row},0)")
    ws.cell(total_row, start_col + 9, f"=IF({get_column_letter(start_col + 10)}{total_row}>0,EXP(-($AF$6/{get_column_letter(start_col + 10)}{total_row})),0)")
    for col_offset in (7, 8, 9):
        ws.cell(total_row, start_col + col_offset).number_format = "0.0%"
    for col_offset in (10, 11):
        ws.cell(total_row, start_col + col_offset).number_format = "0.0"
    dv = DataValidation(type="list", formula1='"Disponible,No Disponible,FUERA"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{get_column_letter(start_col + 12)}{first_data}:{get_column_letter(start_col + 12)}{last_data}")
    return total_row


def _write_helper_and_chart(ws, report: dict[str, Any], table_start: int, table_start_col: int, total_row: int) -> None:
    rows = report.get("rows") if isinstance(report.get("rows"), list) else []
    helper_row = 2
    helper_col = 35  # AI
    headers = ["Eco", "Disp", "Util", "Conf", "TMEF", "TMPR"]
    for idx, header in enumerate(headers):
        cell = ws.cell(helper_row, helper_col + idx, header)
        cell.font = _font(8, True, WHITE)
        cell.fill = _fill(WHITE)
    for i, _row in enumerate(rows, start=3):
        source_row = table_start + i - 2
        ws.cell(i, helper_col, f"={get_column_letter(table_start_col)}{source_row}")
        ws.cell(i, helper_col + 1, f'=IF(ISNUMBER({get_column_letter(table_start_col + 7)}{source_row}),{get_column_letter(table_start_col + 7)}{source_row}*100,0)')
        ws.cell(i, helper_col + 2, f'=IF(ISNUMBER({get_column_letter(table_start_col + 8)}{source_row}),{get_column_letter(table_start_col + 8)}{source_row}*100,0)')
        ws.cell(i, helper_col + 3, f'=IF(ISNUMBER({get_column_letter(table_start_col + 9)}{source_row}),{get_column_letter(table_start_col + 9)}{source_row}*100,0)')
        ws.cell(i, helper_col + 4, f"={get_column_letter(table_start_col + 10)}{source_row}")
        ws.cell(i, helper_col + 5, f"={get_column_letter(table_start_col + 11)}{source_row}")
    if not rows:
        return
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "% Disponibilidad"
    chart.height = 8.4
    chart.width = 13.8
    chart.visible_cells_only = False
    data = Reference(ws, min_col=helper_col + 1, min_row=helper_row, max_row=helper_row + len(rows))
    cats = Reference(ws, min_col=helper_col, min_row=helper_row + 1, max_row=helper_row + len(rows))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 110
    chart.y_axis.numFmt = "0"
    chart.y_axis.majorGridlines = chart.y_axis.majorGridlines
    chart.x_axis.title = "% Disponibilidad"
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    ws.add_chart(chart, "J10")


def _write_dashboard_sheet(wb: Workbook, report: dict[str, Any], settings: dict[str, Any]) -> None:
    sheet_name, title = _group_title(str(report.get("group") or ""))
    ws = wb.create_sheet(sheet_name)
    _configure_sheet(ws)
    _style_range(ws, 1, 1, 60, 31, fill=BG)
    _style_range(ws, 3, 1, 28, 8, fill=WHITE)
    _style_range(ws, 3, 10, 28, 22, fill=WHITE)
    _style_range(ws, 3, 23, 28, 30, fill=WHITE)
    ws["A1"] = "MGA"
    ws["A1"].font = _font(13, True, WHITE)
    ws["A1"].alignment = _align()
    ws["A1"].fill = _fill(BG)
    _style_range(ws, 2, 1, 2, 2, fill=MAGENTA)
    _merge_value(ws, "L1:V1", title, font=_font(21, True, NAVY), fill=BG, border=Border())

    settings = settings or {}
    ws["AF1"] = _setting(settings, "meta_availability", 85) / 100
    ws["AF2"] = _setting(settings, "meta_utilization", 75) / 100
    ws["AF3"] = _setting(settings, "meta_reliability", 80) / 100
    ws["AF4"] = _setting(settings, "meta_tmef", 24)
    ws["AF5"] = _setting(settings, "meta_tmpr", 4)
    ws["AF6"] = _setting(settings, "reliability_mission_hours", _setting(settings, "mission_hours", 24))

    table_start = 35
    table_start_col = 3
    total_row = _write_table(ws, report, table_start, table_start_col, settings)
    total_disp = f"J{total_row}"
    total_util = f"K{total_row}"
    total_conf = f"L{total_row}"
    total_tmef = f"M{total_row}"
    total_tmpr = f"N{total_row}"

    _write_metric_pair(
        ws,
        3,
        14,
        (1, 4),
        (5, 8),
        "% Disponibilidad",
        f"={total_disp}",
        total_disp,
        '="Meta "&TEXT($AF$1,"0.0%")',
        "pink",
        "Meta",
        "=$AF$1",
        "$AF$1",
        f"={total_disp}-$AF$1",
        "white",
    )
    _write_metric_pair(
        ws,
        16,
        28,
        (1, 4),
        (5, 8),
        "% Utilizacion",
        f"={total_util}",
        total_util,
        '="Meta "&TEXT($AF$2,"0.0%")',
        "pink",
        "Meta",
        "=$AF$2",
        "$AF$2",
        f"={total_util}-$AF$2",
        "white",
    )
    _write_metric_pair(
        ws,
        3,
        14,
        (23, 26),
        (27, 30),
        "Confiabilidad",
        f"={total_conf}",
        total_conf,
        '="Meta "&TEXT($AF$3,"0.0%")',
        "white",
        "TMEF",
        f"={total_tmef}",
        f"{total_tmef}/$AF$4",
        '="Meta "&TEXT($AF$4,"0.0")&" h"',
        "white",
        right_number_format="0.0",
    )
    _write_metric_pair(
        ws,
        16,
        28,
        (23, 26),
        (27, 30),
        "TMPR",
        f"={total_tmpr}",
        f"MAX(0,1-({total_tmpr}/$AF$5))",
        '="Meta "&TEXT($AF$5,"0.0")&" h"',
        "pink",
        "Meta Conf.",
        "=$AF$3",
        "$AF$3",
        f"={total_conf}-$AF$3",
        "white",
        left_number_format="0.0",
    )

    tab_border = _border("000000")
    for cell_range, label, fill in (
        ("J4:L6", "%\nDisponibilidad", TEAL),
        ("M4:O6", "% Utilizacion", WHITE),
        ("P4:R6", "Confiabilidad", WHITE),
        ("S4:V6", "TMEF", WHITE),
        ("J7:L9", "TMPR", WHITE),
    ):
        _merge_value(ws, cell_range, label, font=_font(9, False, "00263A" if fill == TEAL else "222222"), fill=fill, border=tab_border)
        ws[cell_range.split(":")[0]].alignment = _align("left", "center", True)

    _write_helper_and_chart(ws, report, table_start + 1, table_start_col, total_row)
    _merge_value(ws, "A30:AD30", "REPORTE SEMANAL DE INDICADORES", font=_font(14, True, "000000"), fill=BG, border=Border())
    _merge_value(ws, "J32:K32", "Dia Inicial:", font=_font(9, True, "7A7A7A"), fill=BG, border=Border())
    _merge_value(ws, "L32:N32", int(report.get("start_day") or str(report.get("start", "01"))[-2:] or 1), font=_font(11, True, "000000"), fill=WHITE, border=Border())
    _merge_value(ws, "P32:Q32", "Dia Final:", font=_font(9, True, "7A7A7A"), fill=BG, border=Border())
    _merge_value(ws, "R32:T32", int(report.get("end_day") or str(report.get("end", "30"))[-2:] or 30), font=_font(11, True, "000000"), fill=WHITE, border=Border())
    _merge_value(ws, "C45:S45", "Editable: modifica horas, estatus o porcentajes de la tabla. La fila Total alimenta el tablero.", font=_font(8, False, "737373"), fill=BG, border=Border())
    for col in range(32, 41):
        ws.column_dimensions[get_column_letter(col)].hidden = True


def build_editable_kpi_excel(reports: list[dict[str, Any]], settings: dict[str, Any] | None = None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for report in reports:
        _write_dashboard_sheet(wb, report, settings or {})
    instructions = wb.create_sheet("INSTRUCCIONES")
    instructions.sheet_state = "hidden"
    instructions["A1"] = "Uso"
    instructions["A1"].font = _font(14, True, NAVY)
    instructions["A3"] = "Editar porcentajes"
    instructions["B3"] = "Modifica la tabla visible y la fila Total; las tarjetas superiores leen la fila Total."
    instructions["A4"] = "Editar metas"
    instructions["B4"] = "Las metas estan ocultas en AF1:AF6 de cada hoja."
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
